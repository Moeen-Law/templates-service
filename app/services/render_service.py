import logging
import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from jinja2 import BaseLoader, Environment

from app.utils.jinja_utils import extract_variables_from_text

logger = logging.getLogger(__name__)

_RLM = "\u200f"


def _set_document_rtl_defaults(doc: Document, font_name: str, body_size: int) -> None:
    settings_el = doc.settings.element
    if settings_el.find(qn("w:bidi")) is None:
        settings_el.insert(0, OxmlElement("w:bidi"))

    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return

    ppr_default_el = doc_defaults.find(".//" + qn("w:pPr"))
    if ppr_default_el is not None and ppr_default_el.find(qn("w:bidi")) is None:
        ppr_default_el.insert(0, OxmlElement("w:bidi"))

    rpr_default_el = doc_defaults.find(".//" + qn("w:rPr"))
    if rpr_default_el is not None:
        lang = rpr_default_el.find(qn("w:lang"))
        if lang is not None:
            lang.set(qn("w:val"), "ar-SA")
            lang.set(qn("w:bidi"), "ar-SA")

        r_fonts = rpr_default_el.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            rpr_default_el.insert(0, r_fonts)
        r_fonts.set(qn("w:cs"), font_name)

        sz_cs = rpr_default_el.find(qn("w:szCs"))
        if sz_cs is None:
            sz_cs = OxmlElement("w:szCs")
            rpr_default_el.append(sz_cs)
        sz_cs.set(qn("w:val"), str(body_size * 2))


def _set_normal_style_rtl(doc: Document, font_name: str, body_size: int) -> None:
    normal = doc.styles["Normal"]
    el = normal.element

    ppr = el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        el.insert(0, ppr)
    if ppr.find(qn("w:bidi")) is None:
        ppr.insert(0, OxmlElement("w:bidi"))

    rpr = el.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        el.append(rpr)
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.insert(0, r_fonts)
    r_fonts.set(qn("w:cs"), font_name)

    sz_cs = rpr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = OxmlElement("w:szCs")
        rpr.append(sz_cs)
    sz_cs.set(qn("w:val"), str(body_size * 2))


def _make_rtl_paragraph(
    doc: Document,
    alignment: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.RIGHT,
):
    para = doc.add_paragraph()
    ppr = para._p.get_or_add_pPr()

    if ppr.find(qn("w:bidi")) is None:
        ppr.insert(0, OxmlElement("w:bidi"))

    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    jc.set(qn("w:val"), "center" if alignment == WD_ALIGN_PARAGRAPH.CENTER else "right")

    para.alignment = alignment
    return para


def _stamp_run(run, font_name: str, size_pt: int) -> None:
    rpr = run._r.get_or_add_rPr()

    if rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))

    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), font_name)

    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:bidi"), "ar-SA")

    run.font.size = Pt(size_pt)


def _with_rtl_mark(text: str) -> str:
    return text if text.strip() == "" else f"{_RLM}{text}"


def _add_inline(para, text: str, font_name: str, size_pt: int, base_bold: bool = False):
    segments = re.split(r"(\*\*[^*]+\*\*)", text)
    for segment in segments:
        if not segment:
            continue
        bold_match = re.fullmatch(r"\*\*([^*]+)\*\*", segment)
        if bold_match:
            run = para.add_run(_with_rtl_mark(bold_match.group(1)))
            run.bold = True
        else:
            run = para.add_run(_with_rtl_mark(segment))
            run.bold = base_bold
        _stamp_run(run, font_name, size_pt)


def _add_rule(doc: Document) -> None:
    para = _make_rtl_paragraph(doc)
    ppr = para._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)
    ppr.append(p_bdr)
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(2)


def _add_two_col_line(doc: Document, line: str, font_name: str, size_pt: int) -> None:
    parts = [part.strip() for part in line.split("|", 1)]
    para = _make_rtl_paragraph(doc)
    _add_inline(para, parts[0], font_name, size_pt)
    tab_run = para.add_run("\t")
    _stamp_run(tab_run, font_name, size_pt)
    if len(parts) > 1:
        _add_inline(para, parts[1], font_name, size_pt)
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)


def _build_document(
    doc: Document,
    markdown: str,
    font_name: str,
    body_size: int,
    h1_size: int,
    h2_size: int,
):
    for line in markdown.splitlines():
        if line.startswith("# "):
            para = _make_rtl_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER)
            run = para.add_run(_with_rtl_mark(line[2:].strip()))
            run.bold = True
            _stamp_run(run, font_name, h1_size)
            para.paragraph_format.space_after = Pt(6)
        elif line.startswith("## "):
            para = _make_rtl_paragraph(doc)
            run = para.add_run(_with_rtl_mark(line[3:].strip()))
            run.bold = True
            run.underline = True
            _stamp_run(run, font_name, h2_size)
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(4)
        elif line.strip() == "---":
            _add_rule(doc)
        elif line.startswith("- "):
            para = _make_rtl_paragraph(doc)
            ppr = para._p.get_or_add_pPr()
            ind = OxmlElement("w:ind")
            ind.set(qn("w:right"), "360")
            ind.set(qn("w:hanging"), "360")
            ppr.append(ind)
            bullet_run = para.add_run(_with_rtl_mark("- "))
            _stamp_run(bullet_run, font_name, body_size)
            _add_inline(para, line[2:].strip(), font_name, body_size)
            para.paragraph_format.space_after = Pt(3)
        elif "|" in line and line.strip():
            _add_two_col_line(doc, line, font_name, body_size)
        elif line.strip() == "":
            para = _make_rtl_paragraph(doc)
            para.paragraph_format.space_after = Pt(2)
        else:
            para = _make_rtl_paragraph(doc)
            _add_inline(para, line, font_name, body_size)
            para.paragraph_format.space_after = Pt(4)


class RenderService:
    font_name = "Traditional Arabic"
    body_size = 12
    h1_size = 18
    h2_size = 14

    @staticmethod
    def extract_placeholders(markdown_content: str) -> set[str]:
        variables = extract_variables_from_text(markdown_content)
        logger.debug("Extracted placeholders count=%s", len(variables))
        return set(variables)

    @classmethod
    def render_docx(cls, markdown_content: str, data: dict[str, Any]) -> bytes:
        logger.debug(
            "Rendering markdown template chars=%s data_keys=%s",
            len(markdown_content),
            sorted(data.keys()),
        )
        env = Environment(
            loader=BaseLoader(),
            trim_blocks=True,
            lstrip_blocks=True,
        )
        rendered_md = env.from_string(markdown_content).render(**data)

        doc = Document()
        section = doc.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

        _set_document_rtl_defaults(doc, cls.font_name, cls.body_size)
        _set_normal_style_rtl(doc, cls.font_name, cls.body_size)
        _build_document(
            doc, rendered_md, cls.font_name, cls.body_size, cls.h1_size, cls.h2_size
        )

        buffer = BytesIO()
        doc.save(buffer)
        rendered = buffer.getvalue()
        logger.debug("Rendered docx size_bytes=%s", len(rendered))
        return rendered
