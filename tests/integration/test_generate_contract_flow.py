from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient

from app.integrations.file_service import FileServiceClient
from app.main import app


def _build_template_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Contract between {{ client_name }} and {{ company_name }}")
    doc.add_paragraph("Start date: {{ start_date }}")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_generate_contract_end_to_end_with_mocked_file_service(monkeypatch):
    template_bytes = _build_template_docx_bytes()

    async def fake_download(self, file_id: str) -> bytes:
        assert file_id == "template-file-1"
        return template_bytes

    async def fake_upload(self, filename: str, content: bytes) -> str:
        assert filename == "Employment Contract.docx"
        assert isinstance(content, bytes)
        assert len(content) > 0
        return "generated-file-123"

    monkeypatch.setattr(FileServiceClient, "download_template", fake_download)
    monkeypatch.setattr(FileServiceClient, "upload_document", fake_upload)

    with TestClient(app) as client:
        create_response = client.post(
            "/templates",
            json={
                "name": "Employment Contract",
                "description": "Basic employment agreement",
                "file_id": "template-file-1",
                "fields": [
                    {"name": "client_name", "type": "string", "required": True},
                    {"name": "company_name", "type": "string", "required": True},
                    {"name": "start_date", "type": "date", "required": True},
                ],
            },
        )
        assert create_response.status_code == 201

        template_id = create_response.json()["id"]

        generate_response = client.post(
            "/contracts/generate",
            json={
                "template_id": template_id,
                "data": {
                    "client_name": "Ahmed",
                    "company_name": "Moeen",
                    "start_date": "2026-04-11",
                },
            },
        )

        assert generate_response.status_code == 200
        assert generate_response.json() == {"file_id": "generated-file-123"}


def test_generate_contract_returns_validation_errors(monkeypatch):
    template_bytes = _build_template_docx_bytes()

    async def fake_download(self, file_id: str) -> bytes:
        return template_bytes

    async def fake_upload(self, filename: str, content: bytes) -> str:
        return "generated-file-123"

    monkeypatch.setattr(FileServiceClient, "download_template", fake_download)
    monkeypatch.setattr(FileServiceClient, "upload_document", fake_upload)

    with TestClient(app) as client:
        create_response = client.post(
            "/templates",
            json={
                "name": "Employment Contract",
                "description": "Basic employment agreement",
                "file_id": "template-file-1",
                "fields": [
                    {"name": "client_name", "type": "string", "required": True},
                    {"name": "company_name", "type": "string", "required": True},
                    {"name": "start_date", "type": "date", "required": True},
                ],
            },
        )
        template_id = create_response.json()["id"]

        generate_response = client.post(
            "/contracts/generate",
            json={
                "template_id": template_id,
                "data": {
                    "client_name": "Ahmed",
                    "start_date": "not-a-date",
                },
            },
        )

        assert generate_response.status_code == 422
        errors = generate_response.json()["errors"]
        assert {item["type"] for item in errors} == {"missing_field", "invalid_type"}


def test_create_template_with_file_upload_flow(monkeypatch):
    template_bytes = _build_template_docx_bytes()

    captured: dict[str, object] = {}

    async def fake_upload_template_file(
        self,
        filename: str,
        content: bytes,
        content_type: str | None = None,
    ) -> str:
        captured["filename"] = filename
        captured["content_type"] = content_type
        captured["content_size"] = len(content)
        return "uploaded-template-file-id"

    monkeypatch.setattr(
        FileServiceClient,
        "upload_template_file",
        fake_upload_template_file,
    )

    with TestClient(app) as client:
        create_response = client.post(
            "/templates",
            data={
                "name": "Employment Contract Uploaded",
                "description": "Uploaded through multipart",
                "fields": json.dumps(
                    [
                        {"name": "client_name", "type": "string", "required": True},
                        {
                            "name": "company_name",
                            "type": "string",
                            "required": True,
                        },
                        {"name": "start_date", "type": "date", "required": True},
                    ]
                ),
            },
            files={
                "file": (
                    "employment_contract.docx",
                    template_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert create_response.status_code == 201
    payload = create_response.json()
    assert payload["file_id"] == "uploaded-template-file-id"
    assert payload["name"] == "Employment Contract Uploaded"
    assert len(payload["fields"]) == 3
    assert captured["filename"] == "employment_contract.docx"
    assert captured["content_size"] == len(template_bytes)
